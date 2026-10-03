"""Cross-check app/src/data/rules.json against a StrategyWiki-derived workbook.

Usage: python tools/compare_wiki.py <KOTOR2_Classes_Feats_Powers.xlsx>

Prints every disagreement; the game files are the source of truth, this only
points at things worth checking in game.
"""
import json
import re
import sys
from collections import defaultdict

import openpyxl

RULES = json.load(open("app/src/data/rules.json", encoding="utf-8"))
CLS = {c["name"]: c for c in RULES["classes"]}
FEAT = {f["name"].lower(): f for f in RULES["feats"]}
POWER = {p["name"].lower(): p for p in RULES["powers"]}
SKILLS = [s["name"] for s in RULES["skills"]]


def rows(ws, header_row):
    head = [c.value for c in ws[header_row]]
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if any(v is not None for v in r):
            yield dict(zip(head, r))


def split(v):
    return [x.strip() for x in str(v).split(";") if x.strip()] if v else []


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = defaultdict(list)

    # Class summary: skills, per-level vitality/force
    for r in rows(wb["Class Summary"], 4):
        c = CLS.get(r["Class"])
        if not c:
            continue
        if r["Vitality / level"] != c["hitDie"]:
            out["class"].append(f"{c['name']}: vitality/level wiki {r['Vitality / level']} vs game {c['hitDie']}")
        if r["Force / level"] != c["forceDie"]:
            out["class"].append(f"{c['name']}: force/level wiki {r['Force / level']} vs game {c['forceDie']}")
        wiki_sk = {s for s in SKILLS if r.get(f"Skill: {s}") == "Yes"}
        game_sk = {SKILLS[i] for i in c["classSkills"]}
        if wiki_sk != game_sk:
            out["class skills"].append(f"{c['name']}: wiki-only {sorted(wiki_sk - game_sk)}, game-only {sorted(game_sk - wiki_sk)}")

    # Per-level progression
    for r in rows(wb["All Levels"], 4):
        c = CLS.get(r["Class"])
        if not c:
            continue
        i = r["Level"] - 1
        tag = f"{c['name']} L{r['Level']}"
        for col, key in (("Fortitude", "fort"), ("Reflex", "ref"), ("Will", "will"), ("Defense bonus", "classDefense")):
            j = i + 1 if key == "classDefense" else i  # acbonus rows start at level 0
            if r[col] is not None and j < len(c[key]) and r[col] != c[key][j]:
                out[col].append(f"{tag}: wiki {r[col]} vs game {c[key][j]}")
        sel = r["Feat selected"]
        if sel in ("Yes", "No") and i < len(c["featPicks"]) and (sel == "Yes") != bool(c["featPicks"][i]):
            out["feat picks"].append(f"{tag}: wiki {sel} vs game {c['featPicks'][i]}")
        ps = r["Powers selected"]
        if isinstance(ps, int) and i < len(c["powerPicks"]) and ps != c["powerPicks"][i]:
            out["power picks"].append(f"{tag}: wiki {ps} vs game {c['powerPicks'][i]}")
        game_feats = {f["name"].lower() for f in RULES["feats"] if f["classes"].get(c["id"], {}).get("grant") == r["Level"]}
        wiki_feats = {x.lower() for x in split(r["Feats granted"])}
        if wiki_feats != game_feats:
            out["granted feats"].append(f"{tag}: wiki-only {sorted(wiki_feats - game_feats)}, game-only {sorted(game_feats - wiki_feats)}")
        game_pw = {next(p["name"] for p in RULES["powers"] if p["id"] == pid).lower() for lv, pid in c["powerGrants"] if lv == r["Level"]}
        wiki_pw = {x.lower() for x in split(r["Powers granted"])}
        if wiki_pw != game_pw:
            out["granted powers"].append(f"{tag}: wiki {sorted(wiki_pw)} vs planner {sorted(game_pw)}")

    # Power unlock levels
    for r in rows(wb["Power Unlocks"], 4):
        p = POWER.get(str(r["Power"]).lower())
        if not p and r["Side"] is None:
            continue  # summary/footnote rows
        if not p:
            out["power unlocks"].append(f"{r['Power']}: not in planner")
            continue
        lv = set(p["classes"].values())
        if lv and r["Min. character level"] not in lv:
            out["power unlocks"].append(f"{p['name']}: wiki L{r['Min. character level']} vs game {sorted(lv)}")
        side = {"Light side": "light", "Dark side": "dark", "Universal": "universal"}.get(r["Side"])
        if side and side != p["side"]:
            out["power side"].append(f"{p['name']}: wiki {side} vs game {p['side']}")

    # Feat prerequisites / min level
    for r in rows(wb["Feats"], 4):
        if r["Attribute"] != "Prerequisites":
            continue
        f = FEAT.get(str(r["Tier"]).lower())
        if not f:
            continue
        m = re.search(r"Level (\d+)", str(r["Value"]))
        if m and int(m.group(1)) != f["minLevel"]:
            out["feat min level"].append(f"{f['name']}: wiki L{m.group(1)} vs game L{f['minLevel']}")

    # Prestige feat restrictions
    pairs = {"Jedi Weapon Master / Sith Marauder": ["jwm", "sma"], "Jedi Watchman / Sith Assassin": ["jwa", "sas"],
             "Jedi Master / Sith Lord": ["jma", "sld"]}
    for r in rows(wb["Prestige Feat Rules"], 4):
        names = re.split(r",\s*", str(r["Feat(s)"]))
        for col, ids in pairs.items():
            if r.get(col) not in ("Yes", "No"):
                continue
            for n in names:
                f = FEAT.get(n.lower())
                if not f:
                    continue
                for cid in ids:
                    game = bool(f["classes"].get(cid, {}).get("pick"))
                    if game != (r[col] == "Yes"):
                        out["prestige feat rules"].append(f"{f['name']} for {cid}: wiki {r[col]} vs game {'pickable' if game else 'not pickable'}")

    for k, v in out.items():
        print(f"\n## {k} ({len(v)})")
        for line in v:
            print("  " + line)
    if not out:
        print("No disagreements.")


if __name__ == "__main__":
    main(sys.argv[1])
